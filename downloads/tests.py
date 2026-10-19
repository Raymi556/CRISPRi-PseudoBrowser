from django.test import TestCase
from io import BytesIO

from openpyxl import load_workbook

from .models import (
    Comparison,
    Gene,
    GeneConditionSummary,
    Guide,
    GuideObservation,
    PathwayRecord,
)


class DownloadModelTests(TestCase):
    def test_core_models_can_be_related(self):
        gene = Gene.objects.create(pao1_id="PA0001", gene_name="test")
        comparison = Comparison.objects.create(
            batch="b1",
            comparison_id="b1_test",
            raw_condition="test",
            display_label="Test",
            comparison_class="control",
            comparison_order=1,
        )
        guide = Guide.objects.create(
            identity_key="PA0001:AAAAAAAAAAAAAAAAAAAA",
            gene=gene,
            spacer="AAAAAAAAAAAAAAAAAAAA",
            guide_type="perfect",
        )
        summary = GeneConditionSummary.objects.create(
            gene=gene,
            comparison=comparison,
            guide_type="perfect",
            med_lfc=0.0,
            p_stouffer=1.0,
        )

        self.assertEqual(guide.gene, gene)
        self.assertEqual(summary.comparison, comparison)

    def test_download_page_lists_pathway_exports(self):
        response = self.client.get("/downloads/")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "COG pathway records")
        self.assertContains(response, "Gene Ontology pathway records")
        self.assertContains(response, "Guide observations")
        self.assertContains(response, "Gene-condition summaries")
        self.assertNotContains(response, "Generated on download")

    def test_pathway_export_is_excel(self):
        PathwayRecord.objects.create(
            dataset="cog",
            pao1_id="PA0001",
            accession="COG0001",
            term_name="Replication",
            category="replication",
        )
        response = self.client.get("/downloads/pathway_cog/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Disposition"],
            'attachment; filename="pao1_cog.xlsx"',
        )
        content = b"".join(response.streaming_content)
        workbook = load_workbook(BytesIO(content), read_only=True)
        self.assertEqual(
            list(workbook.active.values)[0][:2],
            ("pao1_id", "gene_name"),
        )
        workbook.close()

    def test_guide_observation_export_uses_database_rows(self):
        gene = Gene.objects.create(pao1_id="PA0002", gene_name="test_gene")
        comparison = Comparison.objects.create(
            batch="b1",
            comparison_id="b1_test",
            raw_condition="test",
            display_label="Test",
            comparison_class="control",
            comparison_order=1,
        )
        guide = Guide.objects.create(
            identity_key="PA0002:AAAAAAAAAAAAAAAAAAAA",
            gene=gene,
            spacer="AAAAAAAAAAAAAAAAAAAA",
            guide_type="perfect",
        )
        GuideObservation.objects.create(
            guide=guide,
            comparison=comparison,
            lfc_adj=-1.25,
        )

        response = self.client.get("/downloads/guide_observations/")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(
            BytesIO(b"".join(response.streaming_content)), read_only=True
        )
        rows = list(workbook.active.values)
        self.assertIn("guide__gene__pao1_id", rows[0])
        flattened = [value for row in rows for value in row]
        self.assertIn("PA0002", flattened)
        self.assertIn(-1.25, flattened)
        workbook.close()
